"""Create a KangKangPet JSON reminder backup from the renewal workbook."""

from __future__ import annotations

import datetime as dt
import json
import uuid
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path("D:/Codex/域名")
SOURCE = ROOT / "续费清单.xlsx"
TODAY = dt.date(2026, 8, 13)
EXPORTED_AT = "2026-08-13T00:00:00.000Z"
IMPORT_FILE = ROOT / "康康熊提醒一键导入_未来有效.json"
AUDIT_FILE = ROOT / "续费清单_康康熊导入核对.xlsx"


def clean(value):
    return str(value).replace(" ", " ").strip() if value is not None else ""


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(clean(value)[:10])
    except ValueError:
        return None


def info(*pairs):
    return " | ".join(f"{label}: {clean(value)}" for label, value in pairs if clean(value))


def source_records():
    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    records = []
    for project, domain, expires, region, status, department, cost, note, renewal_to in workbook["域名"].iter_rows(min_row=2, values_only=True):
        if clean(domain):
            records.append({"category": "域名", "title": f"【域名续费】{clean(domain)}", "expires": date_value(renewal_to) or date_value(expires), "status": clean(status), "department": clean(department), "cost": clean(cost), "cycle": "按年", "message": info(("区域", region), ("状态", status), ("使用部门", department), ("费用", cost), ("备注", note), ("续费至", renewal_to))})

    group = ""
    for item_group, name, purpose, status, department, expires, cost, cycle, note in workbook["服务器"].iter_rows(min_row=2, values_only=True):
        if clean(item_group):
            group = clean(item_group)
        if clean(name):
            records.append({"category": "服务器", "title": f"【服务器续费】{clean(name)}", "expires": date_value(expires), "status": clean(status), "department": clean(department), "cost": clean(cost), "cycle": clean(cycle), "message": info(("系统/分组", group), ("用途", purpose), ("状态", status), ("使用部门", department), ("费用", cost), ("续费周期", cycle), ("备注", note))})

    for name, purpose, status, department, expires, cost, cycle, note in workbook["软件"].iter_rows(min_row=2, values_only=True):
        if clean(name):
            suffix = f"（{clean(department)}）" if clean(department) else ""
            records.append({"category": "软件/服务", "title": f"【软件续费】{clean(name)}{suffix}", "expires": date_value(expires), "status": clean(status), "department": clean(department), "cost": clean(cost), "cycle": clean(cycle), "message": info(("用途", purpose), ("状态", status), ("使用部门", department), ("费用", cost), ("续费周期", cycle), ("备注", note))})
    return records


def reminder(record, index):
    date_text = record["expires"].isoformat()
    unique_key = f"{record['category']}|{record['title']}|{date_text}|{record['message']}"
    return {"id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"kangkangpet-renewal/{unique_key}")), "title": record["title"], "message": record["message"], "enabled": True, "sortOrder": index, "repeat": "once", "time": "09:00", "date": date_text, "weekdays": [], "intervalMinutes": None, "notificationMode": "both", "createdAt": EXPORTED_AT, "updatedAt": EXPORTED_AT, "lastTriggeredAt": None, "nextTriggerAt": None, "intervalAnchorAt": EXPORTED_AT}


def audit_book(records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "导入结果"
    headers = ["类别", "事项", "到期时间", "导入状态", "不导入原因", "提示方式", "续费周期", "状态", "使用部门", "费用", "备注/补充信息"]
    sheet.append(headers)
    for record in records:
        expires = record["expires"]
        if expires is None:
            result, reason = "未导入", "原清单缺少有效到期日期"
        elif expires < TODAY:
            result, reason = "未导入", "到期日期已过去；康康熊会拒绝一次性过期提醒"
        else:
            result, reason = "已导入", "JSON 导入包：一次性提醒，09:00，桌宠气泡 + 系统通知"
        sheet.append([record["category"], record["title"], expires.isoformat() if expires else "", result, reason, "桌宠气泡 + 系统通知" if result == "已导入" else "", record["cycle"], record["status"], record["department"], record["cost"], record["message"]])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 12, "B": 36, "C": 14, "D": 12, "E": 45, "F": 28, "G": 16, "H": 12, "I": 18, "J": 16, "K": 70}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["原始数据", "续费清单.xlsx 的 域名、服务器、软件 三个工作表"])
    notes.append(["康康熊导入格式", "JSON 备份，兼容当前“提醒管理 → 导入”的 JSON 提醒备份选择器"])
    notes.append(["已导入规则", "仅导入 2026-08-13 及以后、有有效日期的项目；每条设为 09:00 一次性提醒，同时触发桌宠气泡和系统通知"])
    notes.append(["未导入规则", "已过期或缺少有效日期的项目会被当前康康熊导入校验拒绝，保留在导入结果表供核对和更新日期"])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 105
    for cell in notes[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(AUDIT_FILE)


def main():
    records = source_records()
    future = sorted((record for record in records if record["expires"] and record["expires"] >= TODAY), key=lambda record: (record["expires"], record["title"]))
    payload = {"exportedAt": EXPORTED_AT, "reminderFeatureVersion": 2, "source": "续费清单.xlsx（域名、服务器、软件）", "reminders": [reminder(record, index) for index, record in enumerate(future)]}
    IMPORT_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    audit_book(records)
    print(json.dumps({"records": len(records), "imported": len(future), "skipped": len(records) - len(future), "importFile": str(IMPORT_FILE), "auditFile": str(AUDIT_FILE)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
